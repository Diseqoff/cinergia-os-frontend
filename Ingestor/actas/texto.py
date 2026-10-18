"""Lectura de texto de archivos Word y PDF, y utilidades comunes a todos los lectores."""
from __future__ import annotations

import io
import re
import unicodedata
from datetime import date, timedelta
from typing import NamedTuple

from .modelo import ErrorActa

MESES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7,
    "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12,
}
DIAS = {"lunes": 0, "martes": 1, "miercoles": 2, "jueves": 3, "viernes": 4, "sabado": 5, "domingo": 6}
_INVISIBLES = dict.fromkeys(map(ord, "​‌‍⁠﻿­"), None)


def norm(texto: str) -> str:
    """Minúsculas, sin acentos y con espacios simples."""
    t = unicodedata.normalize("NFD", texto or "")
    t = "".join(c for c in t if unicodedata.category(c) != "Mn").lower()
    return re.sub(r"\s+", " ", t).strip()


def limpiar(texto: str | None) -> str:
    t = (texto or "").translate(_INVISIBLES).replace("\xa0", " ")
    return re.sub(r"[ \t]+", " ", t).strip()


class Documento:
    """Contenido de un archivo en orden de lectura.

    bloques: ("p", texto) para párrafos y ("fila", [[párrafos de la celda], ...]) para filas de tabla.
    """

    def __init__(self, bloques: list[tuple], formato: str):
        self.bloques = bloques
        self.formato = formato

    @property
    def lineas(self) -> list[str]:
        salida: list[str] = []
        for tipo, valor in self.bloques:
            if tipo == "p":
                salida.append(valor)
            else:
                for celda in valor:
                    salida.extend(celda)
        return salida

    @property
    def texto(self) -> str:
        return "\n".join(self.lineas)

    def filas(self) -> list[list[str]]:
        """Filas de tabla con el texto de cada celda unido en una línea."""
        return [[" ".join(c) for c in v] for t, v in self.bloques if t == "fila"]


def _partes(texto: str) -> list[str]:
    """Un párrafo con saltos de línea internos se trata como varias líneas."""
    return [t for t in (limpiar(x) for x in (texto or "").replace("\v", "\n").split("\n")) if t]


def _leer_docx(contenido: bytes) -> Documento:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(contenido))
    bloques: list[tuple] = []
    for el in doc.element.body.iterchildren():
        if el.tag.endswith("}p"):
            for t in _partes(Paragraph(el, doc).text):
                bloques.append(("p", t))
        elif el.tag.endswith("}tbl"):
            for fila in Table(el, doc).rows:
                celdas, previa = [], None
                for c in fila.cells:
                    if c._tc is previa:
                        continue
                    previa = c._tc
                    celdas.append([t for p in c.paragraphs for t in _partes(p.text)])
                if any(celdas):
                    bloques.append(("fila", celdas))
    return Documento(bloques, "docx")


def _leer_pdf(contenido: bytes) -> Documento:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(contenido)
    bloques: list[tuple] = []
    for pagina in pdf:
        texto = pagina.get_textpage().get_text_range()
        for linea in texto.splitlines():
            t = limpiar(linea)
            if t:
                bloques.append(("p", t))
    return Documento(bloques, "pdf")


def leer_documento(nombre_archivo: str, contenido: bytes) -> Documento:
    nombre = nombre_archivo.lower()
    try:
        if nombre.endswith(".docx"):
            doc = _leer_docx(contenido)
        elif nombre.endswith(".pdf"):
            doc = _leer_pdf(contenido)
        else:
            raise ErrorActa("Formato inválido. Solo se admiten archivos .docx o .pdf.")
    except ErrorActa:
        raise
    except Exception:
        raise ErrorActa("No se pudo leer el archivo. Verifica que no esté dañado ni protegido con contraseña.")
    if len(doc.lineas) < 5:
        raise ErrorActa("El archivo no tiene texto legible. Si es un escaneo, súbelo como Word o PDF con texto.")
    return doc


# ---------------------------------------------------------------------------
# Fechas
# ---------------------------------------------------------------------------
_MES_RE = "|".join(MESES)
_FECHA_NUM = re.compile(r"(?<!\d)(\d{1,2})[/-](\d{1,2})[/-](\d{4})(?!\d)")
_RANGO = re.compile(rf"(\d{{1,2}})\s*(?:al|a)\s*(\d{{1,2}})\s*(?:de\s+)?({_MES_RE})(?:\s*(?:de|del)?\s*(\d{{4}}))?")
_LARGA = re.compile(rf"(\d{{1,2}})\s*(?:dias\s+)?(?:de\s+|del\s+mes\s+de\s+)?({_MES_RE})\.?(?:\s*(?:de|del)?\s*(\d{{4}}))?")


class Fecha(NamedTuple):
    fecha: date
    sin_anio: bool = False         # el año se dedujo
    dia_coincide: bool | None = None  # el día de la semana escrito coincide con la fecha


def _crear(a: int, m: int, d: int) -> date | None:
    try:
        return date(a, m, d)
    except ValueError:
        return None


def _dia_escrito(previo: str) -> int | None:
    for nombre, numero in DIAS.items():
        if re.search(rf"\b{nombre}\b", previo):
            return numero
    return None


def _coincide(d: date | None, previo: str) -> bool | None:
    escrito = _dia_escrito(previo)
    return None if d is None or escrito is None else d.weekday() == escrito


def _resolver_sin_anio(dia: int, mes: int, previo: str, referencia: date | None) -> Fecha | None:
    escrito = _dia_escrito(previo)
    if referencia:
        base = _crear(referencia.year, mes, dia)
        if base is None:
            return None
        if base < referencia - timedelta(days=60):
            base = _crear(referencia.year + 1, mes, dia) or base
        if escrito is not None and base.weekday() != escrito:
            alterna = _crear(base.year + 1, mes, dia)
            if alterna and alterna.weekday() == escrito and (alterna - referencia).days <= 200:
                return Fecha(alterna, True, True)
            return Fecha(base, True, False)
        return Fecha(base, True, True if escrito is not None else None)
    if escrito is None:
        return None
    hoy = date.today()
    candidatas = [d for a in range(hoy.year - 1, hoy.year + 3) if (d := _crear(a, mes, dia)) and d.weekday() == escrito]
    if not candidatas:
        return None
    return Fecha(min(candidatas, key=lambda d: abs((d - hoy).days)), True, True)


def fechas_en(texto: str, referencia: date | None = None) -> list[Fecha]:
    """Fechas del texto en orden de aparición.

    Reconoce dd/mm/aaaa, "5 de septiembre de 2026" y rangos como "08 al 14 de junio de 2026"
    (entrega ambos extremos). Si falta el año se deduce con la fecha de referencia, normalmente
    la del acta, o con el día de la semana escrito.
    """
    t = norm(texto)
    halladas: list[tuple[int, Fecha]] = []
    cubiertos: list[tuple[int, int]] = []
    for m in _FECHA_NUM.finditer(t):
        d = _crear(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        if d:
            halladas.append((m.start(), Fecha(d)))
            cubiertos.append((m.start(), m.end()))
    for m in _RANGO.finditer(t):
        mes = MESES[m.group(3)]
        previo = t[max(0, m.start() - 20):m.start()]
        for dia in (m.group(1), m.group(2)):
            if m.group(4):
                d = _crear(int(m.group(4)), mes, int(dia))
                f = Fecha(d, False, _coincide(d, previo)) if d else None
            else:
                f = _resolver_sin_anio(int(dia), mes, previo, referencia)
            if f:
                halladas.append((m.start(), f))
        cubiertos.append((m.start(), m.end()))
    for m in _LARGA.finditer(t):
        if any(a <= m.start() < b for a, b in cubiertos):
            continue
        mes, dia = MESES[m.group(2)], int(m.group(1))
        if m.group(3):
            d = _crear(int(m.group(3)), mes, dia)
            f = Fecha(d, False, _coincide(d, t[max(0, m.start() - 20):m.start()])) if d else None
        else:
            f = _resolver_sin_anio(dia, mes, t[max(0, m.start() - 20):m.start()], referencia)
        if f:
            halladas.append((m.start(), f))
    halladas.sort(key=lambda x: x[0])
    return [f for _, f in halladas]


def primera_fecha(texto: str, referencia: date | None = None) -> Fecha | None:
    f = fechas_en(texto, referencia)
    return f[0] if f else None


def fecha_legible(d: date | None) -> str | None:
    if not d:
        return None
    meses = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
    return f"{d.day:02d} {meses[d.month - 1]} {d.year}"


# ---------------------------------------------------------------------------
# Secciones y etiquetas
# ---------------------------------------------------------------------------
_NUMERADO = re.compile(r"^\s*\d{1,2}\s*[.)]\s+[^\d\s]")
_PREFIJO_NUM = re.compile(r"^\s*(?:\d{1,2}|[ivxIVX]{1,4})\s*[.)\-–]\s*")


def titulo_seccion(linea: str) -> str:
    """Normaliza una línea para compararla con los nombres de sección."""
    t = _PREFIJO_NUM.sub("", linea.strip())
    t = re.sub(r"\(.*?\)", "", t)
    t = norm(t).rstrip(": ").strip()
    return t


def _es_titulo(t: str, prefijo: str) -> bool:
    """Un prefijo terminado en $ debe coincidir exactamente con el título."""
    if prefijo.endswith("$"):
        return t == prefijo[:-1]
    return t == prefijo or (t.startswith(prefijo + " ") and len(t) <= len(prefijo) + 30)


def dividir_secciones(
    lineas: list[str], claves: dict[str, tuple[str, ...]], numerados_desconocidos: bool = True
) -> dict[str, list[str]]:
    """Agrupa las líneas bajo la sección conocida más reciente.

    claves: nombre interno -> prefijos que identifican el título de esa sección.
    Los títulos numerados que no están en la lista abren una sección "otro" para que no se
    mezclen con la anterior. Un título conocido sin número solo cuenta si la línea es corta.
    """
    secciones: dict[str, list[str]] = {"inicio": []}
    actual = "inicio"
    for linea in lineas:
        t = titulo_seccion(linea)
        hallada = None
        if t and len(linea) <= 90:
            for nombre, prefijos in claves.items():
                if any(_es_titulo(t, p) for p in prefijos):
                    hallada = nombre
                    break
        if hallada:
            actual = hallada
            secciones.setdefault(actual, [])
            resto = linea.split(":", 1)[1].strip() if ":" in linea else ""
            if resto:
                secciones[actual].append(resto)
            continue
        if numerados_desconocidos and _NUMERADO.match(linea) and len(linea) <= 80 and not linea.rstrip().endswith((".", ";", ":", ",")):
            actual = "otro"
            secciones.setdefault(actual, [])
            continue
        secciones.setdefault(actual, []).append(linea)
    return secciones


def valor_etiqueta(lineas: list[str], *etiquetas: str, siguiente: bool = True) -> str:
    """Valor de "Etiqueta: valor". Si la etiqueta está sola en su línea, toma la línea siguiente."""
    objetivos = [norm(e) for e in etiquetas]
    for i, linea in enumerate(lineas):
        t = norm(linea)
        for e in objetivos:
            if t == e or t == e + ":":
                if siguiente and i + 1 < len(lineas):
                    return limpiar(lineas[i + 1])
                return ""
            if t.startswith(e) and re.match(r"\s*[:\-–]", t[len(e):]) is not None:
                _, _, valor = linea.partition(":")
                if not _:
                    continue
                return limpiar(valor)
    return ""


_VACIOS = {"", "por definir", "por confirmar", "n/a", "no aplica", "pendiente", "tbd", "-", "—", "por asignar"}


def es_vacio(valor: str | None) -> bool:
    t = norm(valor or "")
    if t in _VACIOS:
        return True
    return bool(re.fullmatch(r"\(?\s*por (definir|confirmar)\s*\)?", t)) or t.startswith(("[", "dd/mm", "escribir ", "elegir "))


def entero_de(valor: str | None) -> int | None:
    """Aforo: número seguido de personas, estudiantes, asistentes, etc. En un rango toma el mayor."""
    if not valor:
        return None
    t = norm(valor)
    rango = re.search(r"(\d{1,4})\s*(?:-|–|a|y)\s*(\d{1,4})\s*(?:personas|estudiantes|asistentes|participantes|alumnos|integrantes|invitados|inscritos)", t)
    if rango:
        return max(int(rango.group(1)), int(rango.group(2)))
    simple = re.search(r"(\d{1,4})\s*(?:personas|estudiantes|asistentes|participantes|alumnos|integrantes|invitados|inscritos|cupos|vacantes)", t)
    if simple:
        return int(simple.group(1))
    solo = re.fullmatch(r"[~≈]?\s*(\d{1,4})\s*", t)
    return int(solo.group(1)) if solo else None


# ---------------------------------------------------------------------------
# Ayudas comunes
# ---------------------------------------------------------------------------
_CORREO = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_VIÑETA = re.compile(r"^\s*(?:[-•●▪◦*·❖➢✓✔►→]\s+|\d{1,2}[.)]\s+|[a-z][.)]\s+|\d+\.\d+(?:\.\d+)*\.?\s+)")


def sin_viñeta(linea: str) -> str:
    return _VIÑETA.sub("", linea.strip()).strip()


def sin_datos_personales(texto: str) -> str:
    """Quita correos y números largos (DNI, celulares) que no deben guardarse."""
    t = _CORREO.sub("", texto)
    t = re.sub(r"\b\d{3}[ ]?\d{3}[ ]?\d{3}\b|\b\d{8,}\b", "", t)
    return re.sub(r"\s{2,}", " ", t).strip(" ,;:")


def recortar(texto: str, largo: int) -> str:
    t = limpiar(texto)
    if len(t) <= largo:
        return t
    corte = t[:largo].rsplit(" ", 1)[0].rstrip(",;:.")
    return corte + "…"


def bloque_etiqueta(lineas: list[str], etiquetas: tuple[str, ...], fines: tuple[str, ...]) -> str:
    """Texto que sigue a una etiqueta hasta la siguiente etiqueta de `fines` (ambas sin acentos)."""
    inicio = None
    for i, linea in enumerate(lineas):
        t = norm(linea)
        for e in etiquetas:
            if t == e or t.startswith(e + ":") or t.startswith(e + " :"):
                inicio = (i, linea.partition(":")[2].strip())
                break
        if inicio:
            break
    if not inicio:
        return ""
    i, resto = inicio
    partes = [resto] if resto else []
    for linea in lineas[i + 1:]:
        t = norm(linea)
        if any(t == f or t.startswith(f + ":") or t.startswith(f + " :") for f in fines):
            break
        partes.append(linea)
    return limpiar(" ".join(partes))


def nombre_persona(linea: str) -> str:
    """Nombre limpio a partir de líneas como "Diego Sánchez (Director de Eventos)" o "Chris: Vicepresidente"."""
    t = sin_viñeta(linea)
    t = re.sub(r"\(.*?\)", "", t)
    if ":" in t:
        antes, _, despues = t.partition(":")
        t = antes if len(antes) <= 45 else despues
    t = re.split(r"\s+[-–—]\s+", t)[0]
    t = re.sub(r"\s+y\s+.*$", "", t) if len(t) > 40 else t
    return t.strip(" .,;")


def nombre_directo(apellidos_nombres: str) -> str:
    """"Agurto Villafuerte, Joaquín Erick" pasa a "Joaquín Erick Agurto Villafuerte"."""
    t = limpiar(re.sub(r"\(.*?\)", "", apellidos_nombres)).strip(" .,;")
    if "," in t:
        ap, _, no = t.partition(",")
        return f"{no.strip()} {ap.strip()}".strip()
    return t
